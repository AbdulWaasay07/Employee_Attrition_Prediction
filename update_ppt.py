import pptx
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

def update_presentation():
    prs = pptx.Presentation('finalppt-1.pptx')
    
    # -------------------------------------------------------------
    # SLIDE 2: Contents
    # -------------------------------------------------------------
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.has_text_frame and '01' in s.text_frame.text:
            s.text_frame.text = (
                "01  Abstract\n"
                "02  Introduction\n"
                "03  Problem Statement\n"
                "04  Objectives & Expected Outcomes\n"
                "05  Literature Review\n"
                "06  Methodology (Phase-1 Completed vs Phase-2)\n"
                "07  Existing System – Open Issues and Challenges\n"
                "08  Proposed System – Advantages and Limitations\n"
                "09  Novelty of the Project\n"
                "10  Phase-1 Completed Progress & Work for Next Presentation\n"
                "11  Major Project Schedule – Gantt Chart\n"
                "12  References"
            )
            for p in s.text_frame.paragraphs:
                p.font.size = Pt(14)
                p.font.name = "Arial"
    
    # -------------------------------------------------------------
    # SLIDE 3: Abstract
    # -------------------------------------------------------------
    slide3 = prs.slides[2]
    # Shape 4: Study focus text
    slide3.shapes[4].text_frame.text = (
        "• Predict employee attrition and diagnose organizational friction using machine learning & XAI on multi-source HR data.\n"
        "• Benchmarking ensemble classifiers (Random Forest & XGBoost) and applying SHAP for directional feature attribution.\n"
        "• Phase-1 Completed: Implemented 6-touchpoint data ingestion, automated cleaning, outlier Winsorization, and 24-D feature store engineering."
    )
    # Shape 8: Base paper evidence
    slide3.shapes[8].text_frame.text = (
        "• Base Dataset: 5,767 employee records with 309 voluntary resignations (5.4% attrition rate).\n"
        "• 27 workforce determinants categorized across Ishikawa root-cause dimensions; 14 core features selected.\n"
        "• Random Forest achieved AUC of 99.9% and 94.7% accuracy with ROSE class-balancing."
    )
    # Shape 9: Bottom banner
    slide3.shapes[9].text_frame.text = (
        "Core contribution: Moving beyond “which features matter?” to “how and in which direction do they drive attrition?”, mapped to actionable HR retention plans."
    )

    # -------------------------------------------------------------
    # SLIDE 6: Objectives & Expected Outcomes
    # -------------------------------------------------------------
    slide6 = prs.slides[5]
    # Move Objectives Title (Shape 0) & Bullets (Shape 2) inside Left Rounded Rectangle (Shape 5)
    slide6.shapes[0].left = Emu(3800000)
    slide6.shapes[0].top = Emu(3200000)
    slide6.shapes[0].width = Emu(5000000)
    slide6.shapes[0].height = Emu(400000)
    slide6.shapes[0].text_frame.text = "Project Objectives"
    for p in slide6.shapes[0].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide6.shapes[2].left = Emu(3800000)
    slide6.shapes[2].top = Emu(3650000)
    slide6.shapes[2].width = Emu(5000000)
    slide6.shapes[2].height = Emu(3400000)
    slide6.shapes[2].text_frame.text = (
        "• Ingest & structure HR data across 6 operational touchpoints into a unified relational database. [Done - Phase 1]\n"
        "• Implement automated data cleaning, missing-value imputation, and Winsorization for salary/workload outliers. [Done - Phase 1]\n"
        "• Engineer a centralized 24-dimensional Feature Store aggregating tenure, compensation, workload, and performance. [Done - Phase 1]\n"
        "• Train & benchmark tree-based ensemble models (Random Forest, XGBoost) using 10-fold cross-validation and imbalance handling.\n"
        "• Apply SHAP Explainable AI (XAI) to quantify positive and negative directional feature influence.\n"
        "• Build an unsupervised K-Means risk cohort segmentation model and deterministic HR retention rules engine."
    )
    for p in slide6.shapes[2].text_frame.paragraphs:
        p.font.size = Pt(11)

    # Move Expected Outcomes Title (Shape 1) & Bullets (Shape 3) inside Right Rounded Rectangle (Shape 6)
    slide6.shapes[1].left = Emu(9500000)
    slide6.shapes[1].top = Emu(3200000)
    slide6.shapes[1].width = Emu(5000000)
    slide6.shapes[1].height = Emu(400000)
    slide6.shapes[1].text_frame.text = "Expected Outcomes"
    for p in slide6.shapes[1].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide6.shapes[3].left = Emu(9500000)
    slide6.shapes[3].top = Emu(3650000)
    slide6.shapes[3].width = Emu(5000000)
    slide6.shapes[3].height = Emu(3400000)
    slide6.shapes[3].text_frame.text = (
        "• A production-ready, standardized 24-D HR feature foundation free of anomalies and extreme outliers. [Delivered - Phase 1]\n"
        "• A high-precision attrition classification model outputting calibrated continuous flight-risk probabilities (0.0 to 1.0).\n"
        "• Directional explainability reports for HR managers, overcoming traditional black-box limitations.\n"
        "• Actionable employee persona segmentation (High Performers, Burnout/Flight Risk, Underutilized).\n"
        "• A full-stack interactive web platform (FastAPI + React) providing automated retention recommendations and stay-interview triggers."
    )
    for p in slide6.shapes[3].text_frame.paragraphs:
        p.font.size = Pt(11)

    # -------------------------------------------------------------
    # SLIDE 8: Methodology
    # -------------------------------------------------------------
    slide8 = prs.slides[7]
    # Update methodology workflow labels
    slide8.shapes[4].text_frame.text = "1. Problem\nDefinition [✔]"
    slide8.shapes[9].text_frame.text = "2. Data Ingestion\n(6 Tables) [✔]"
    slide8.shapes[14].text_frame.text = "3. Cleaning &\n24-D Store [✔]"
    slide8.shapes[19].text_frame.text = "4. Ensemble\nModeling"
    slide8.shapes[24].text_frame.text = "5. Evaluation\n& Selection"
    slide8.shapes[29].text_frame.text = "6. SHAP &\nRetention"

    # Shape 32 & 33: Left Box (Data Ingestion & Preprocessing - Phase 1 Completed)
    slide8.shapes[32].text_frame.text = "Data Ingestion & Preprocessing [Phase-1 Completed ✔]"
    slide8.shapes[33].text_frame.text = (
        "• Base benchmark: 5,767 observations, 309 voluntary resignations across 14 selected determinants.\n"
        "• Platform implementation: Ingested 6 relational HR touchpoints into MySQL (Employees, Compensation, Performance, Workload, HR Tickets, Training).\n"
        "• Preprocessing & Engineering: Implemented schema coercion, missing value imputation, and Winsorization for salary/overtime extremes into a unified 24-D feature store."
    )

    # Shape 36 & 37: Right Box (Modeling, Explainability & Retention - Phase-2 In Progress)
    slide8.shapes[36].text_frame.text = "Modeling, Explainability & Retention [Phase-2 In Progress]"
    slide8.shapes[37].text_frame.text = (
        "• Comparative modeling: Benchmark Random Forest, XGBoost, Decision Trees, and baseline classifiers.\n"
        "• Class imbalance mitigation: ROSE, SMOTE, and boundary sampling with 80:20 train-test splits and 10-fold CV.\n"
        "• XAI & Prescriptive Action: Compute SHAP Shapley values to provide global workforce trends and individual-level directional explanations, driving a deterministic HR retention rules engine."
    )

    # -------------------------------------------------------------
    # SLIDE 10: Proposed System – Advantages and Limitations
    # -------------------------------------------------------------
    slide10 = prs.slides[9]
    # Pipeline box labels
    slide10.shapes[3].text_frame.text = "HR Ingestion\n(6 Tables) [✔]"
    slide10.shapes[7].text_frame.text = "Clean +\nWinsorize [✔]"
    slide10.shapes[11].text_frame.text = "24-D Store\n& Balance"
    slide10.shapes[15].text_frame.text = "RF + XGBoost\nEnsemble"
    slide10.shapes[19].text_frame.text = "SHAP\nXAI"
    slide10.shapes[23].text_frame.text = "Retention\nEngine"

    # Advantages (Shape 27)
    slide10.shapes[27].text_frame.text = (
        "• Unified Data Foundation: Flattens fragmented HR operational touchpoints into an automated, clean 24-D feature store.\n"
        "• Robust Anomaly Handling: Winsorization prevents distorted decision boundaries from extreme salary and workload outliers.\n"
        "• Direction-Aware Explainability: SHAP values reveal whether a variable increases or decreases voluntary flight risk.\n"
        "• Actionable Prescriptive Intelligence: Translates predictive probabilities into deterministic, prioritized HR retention action plans."
    )
    # Limitations (Shape 31)
    slide10.shapes[31].text_frame.text = (
        "• Data Dependency: Predictive efficacy relies on the completeness, consistency, and logging discipline of operational HR data.\n"
        "• Cross-Sectional Dynamics: Static snapshots require periodic retraining to capture evolving organizational and economic shifts.\n"
        "• Human-in-the-Loop Requirement: Model predictions provide decision support and must be evaluated alongside managerial empathy."
    )

    # -------------------------------------------------------------
    # SLIDE 11: Novelty of the Project
    # -------------------------------------------------------------
    slide11 = prs.slides[10]
    # On slide 11:
    # Card 1: Shape 24 (left=3538727, top=2593988, w=3657601, h=1575173)
    # Card 2: Shape 25 (left=7315200, top=2618050, w=3657600, h=1508761)
    # Card 3: Shape 26 (left=11091671, top=2618050, w=3657601, h=1508761)
    # Box 4: Shape 27 (left=3538727, top=4355410, w=5166361, h=2331721)

    # Card 1 Title (Shape 3) & Text (Shape 4)
    slide11.shapes[3].left = Emu(3650000)
    slide11.shapes[3].top = Emu(2700000)
    slide11.shapes[3].width = Emu(3450000)
    slide11.shapes[3].height = Emu(350000)
    slide11.shapes[3].text_frame.text = "1. Direction-Aware Explainability"
    for p in slide11.shapes[3].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide11.shapes[4].left = Emu(3650000)
    slide11.shapes[4].top = Emu(3080000)
    slide11.shapes[4].width = Emu(3450000)
    slide11.shapes[4].height = Emu(1000000)
    slide11.shapes[4].text_frame.text = "• Moves beyond traditional feature rankings to explain how and in which direction (positive/negative force) each variable drives attrition."
    for p in slide11.shapes[4].text_frame.paragraphs:
        p.font.size = Pt(10)

    # Card 2 Title (Shape 7) & Text (Shape 8)
    slide11.shapes[7].left = Emu(7420000)
    slide11.shapes[7].top = Emu(2700000)
    slide11.shapes[7].width = Emu(3450000)
    slide11.shapes[7].height = Emu(350000)
    slide11.shapes[7].text_frame.text = "2. Hybrid Predictive & Cohort ML"
    for p in slide11.shapes[7].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide11.shapes[8].left = Emu(7420000)
    slide11.shapes[8].top = Emu(3080000)
    slide11.shapes[8].width = Emu(3450000)
    slide11.shapes[8].height = Emu(1000000)
    slide11.shapes[8].text_frame.text = "• Pairs supervised ensemble classification (Random Forest/XGBoost) with unsupervised K-Means risk persona clustering (k=3) to profile burnout vs high performers."
    for p in slide11.shapes[8].text_frame.paragraphs:
        p.font.size = Pt(10)

    # Card 3 Title (Shape 11) & Text (Shape 12)
    slide11.shapes[11].left = Emu(11200000)
    slide11.shapes[11].top = Emu(2700000)
    slide11.shapes[11].width = Emu(3450000)
    slide11.shapes[11].height = Emu(350000)
    slide11.shapes[11].text_frame.text = "3. Prescriptive Retention Rules"
    for p in slide11.shapes[11].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide11.shapes[12].left = Emu(11200000)
    slide11.shapes[12].top = Emu(3080000)
    slide11.shapes[12].width = Emu(3450000)
    slide11.shapes[12].height = Emu(1000000)
    slide11.shapes[12].text_frame.text = "• Deterministically maps flight-risk scores and friction indicators into prioritized, actionable HR retention packages (stay interviews, workload rebalancing)."
    for p in slide11.shapes[12].text_frame.paragraphs:
        p.font.size = Pt(10)

    # Box 4 Title (Shape 13) & Text (Shape 14)
    slide11.shapes[13].left = Emu(3650000)
    slide11.shapes[13].top = Emu(4450000)
    slide11.shapes[13].width = Emu(4900000)
    slide11.shapes[13].height = Emu(350000)
    slide11.shapes[13].text_frame.text = "Why It Matters for Modern HR Operations"
    for p in slide11.shapes[13].text_frame.paragraphs:
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(26, 54, 93)

    slide11.shapes[14].left = Emu(3650000)
    slide11.shapes[14].top = Emu(4850000)
    slide11.shapes[14].width = Emu(4900000)
    slide11.shapes[14].height = Emu(1750000)
    slide11.shapes[14].text_frame.text = (
        "• Solves the “black-box” dilemma in HR analytics by providing transparent, defensible reasons for flight-risk alerts.\n"
        "• Uncovers counterintuitive friction patterns (e.g., unexpected bonus perceptions or stagnant promotion velocity) via local SHAP attribution.\n"
        "• Bridges the gap between predictive data science and operational talent management through automated stay-interview workflows."
    )
    for p in slide11.shapes[14].text_frame.paragraphs:
        p.font.size = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 12: Work to be Shown in the Next Presentation
    # -------------------------------------------------------------
    slide12 = prs.slides[11]
    slide12.shapes[0].text_frame.text = "Phase-1 Completed Milestones & Phase-2 Deliverables"

    # Card 1: Data Ingestion (Done)
    slide12.shapes[3].text_frame.text = "1. Data Ingestion [Done ✔]"
    slide12.shapes[4].text_frame.text = "• Ingested 6 relational HR tables with schema validation & referential integrity."
    
    # Card 2: Cleaning & Preprocessing (Done)
    slide12.shapes[7].text_frame.text = "2. Cleaning & 24-D Store [Done ✔]"
    slide12.shapes[8].text_frame.text = "• Applied Winsorization, missing-value imputation, and flattened 6 tables into 24-D Feature Store."

    # Card 3: Ensemble Model Training
    slide12.shapes[11].text_frame.text = "3. Ensemble Model Training"
    slide12.shapes[12].text_frame.text = "• Train & fine-tune Random Forest, XGBoost, and Decision Trees with 10-fold CV & ROSE balancing."

    # Card 4: Multi-Metric Evaluation
    slide12.shapes[15].text_frame.text = "4. Multi-Metric Evaluation"
    slide12.shapes[16].text_frame.text = "• Benchmark ROC-AUC, Sensitivity, Specificity, and F1-scores across models to select the top classifier."

    # Card 5: SHAP Directional Explainability
    slide12.shapes[19].text_frame.text = "5. SHAP Directional XAI"
    slide12.shapes[20].text_frame.text = "• Generate global summary plots and individual waterfall/force plots for local feature attribution."

    # Card 6: Full-Stack Retention Dashboard
    slide12.shapes[23].text_frame.text = "6. Retention Dashboard Demo"
    slide12.shapes[24].text_frame.text = "• Deploy K-Means persona clustering (k=3) and live deterministic retention rules in React + FastAPI."

    # Bottom Banner
    slide12.shapes[25].text_frame.text = "Deliverable: Live demonstration of an end-to-end explainable attrition prediction and proactive retention platform."

    # -------------------------------------------------------------
    # SLIDE 13: Gantt Chart
    # -------------------------------------------------------------
    slide13 = prs.slides[12]
    # Update Task Labels on Gantt Chart to reflect Phase-1 Completed vs Phase-2
    slide13.shapes[99].text_frame.text = "Literature & Problem Definition [Done ✔]"
    slide13.shapes[109].text_frame.text = "Multi-Source Data Ingestion [Done ✔]"
    slide13.shapes[120].text_frame.text = "Cleaning & 24-D Feature Store [Done ✔]"
    slide13.shapes[132].text_frame.text = "Ensemble Modeling & Balancing [Phase-2]"
    slide13.shapes[143].text_frame.text = "SHAP & Directional XAI [Phase-2]"
    slide13.shapes[154].text_frame.text = "Retention Rules & Dashboard [Phase-2]"
    slide13.shapes[165].text_frame.text = "Testing, Docs & Final Presentation [Phase-2]"

    prs.save('finalppt-1.pptx')
    print('Successfully updated finalppt-1.pptx!')

if __name__ == '__main__':
    update_presentation()
